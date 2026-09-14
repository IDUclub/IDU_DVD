import pytest

from src.dvd_service.modules.doc_parsers import DocumentParser
from src.dvd_service.modules.hierarchy import HierarchyBuilder
from src.dvd_service.modules.structure import StructureTagger


@pytest.mark.parametrize(
    "text",
    [
        "(часть 3.3 в ред. Федерального закона от 03.08.2018 N 342-ФЗ)",
        "3.2. Работы выполняют в соответствии с частью 3.3 настоящей статьи.",
    ],
)
def test_inline_references_and_dates_are_not_new_parts(settings, text):
    settings.split_sentences = False
    assert DocumentParser(settings)._split_block(text) == [text]


def test_inserted_parts_belong_to_article_not_part_three():
    parts = [
        {"id": i, "numbering": num, "type": typ, "text": txt, "relation": "deeper"}
        for i, (num, typ, txt) in enumerate(
            [
                ("6", "chapter", "Глава 6. Строительство"),
                ("52", "article", "Статья 52. Строительство"),
                ("3", "clause", "Лицо, осуществляющее строительство"),
                ("3.3", "clause", "По решению застройщика"),
                ("53", "article", "Статья 53. Строительный контроль"),
                ("3.3", "clause", "Иное положение"),
            ]
        )
    ]
    hb = HierarchyBuilder()
    nodes = hb.flatten(hb.build(parts, StructureTagger(None).numbering_ranks(parts)))
    by_id = {n["id"]: n for n in nodes}
    found = [n for n in nodes if n["numbering"] == "3.3"]
    assert [by_id[n["parent_id"]]["numbering"] for n in found] == ["52", "53"]


def make_word_fixture(path):
    from docx import Document
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

    doc = Document()
    numbering = doc.part.numbering_part.element
    numbering.append(parse_xml(f"""<w:abstractNum {nsdecls('w')} w:abstractNumId="900">
      <w:lvl w:ilvl="0"><w:start w:val="3"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
      <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
    </w:abstractNum>"""))
    numbering.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="900"><w:abstractNumId w:val="900"/></w:num>'
        )
    )
    doc.add_paragraph("Глава 6. Строительство")
    doc.add_paragraph("Статья 52. Осуществление строительства")
    texts = [
        (0, "Лицо, осуществляющее строительство."),
        (1, "Застройщик вправе выполнять работы."),
        (1, "Работы выполняют в соответствии с частью 3.3 настоящей статьи."),
        (
            1,
            "По решению застройщика этапы строительства могут быть выделены после получения разрешения.",
        ),
    ]
    for level, text in texts:
        p = doc.add_paragraph(text)
        num = p._p.get_or_add_pPr().get_or_add_numPr()
        num.get_or_add_numId().val = 900
        num.get_or_add_ilvl().val = level
        if level:
            doc.add_paragraph("(в ред. Федерального закона от 27.12.2019 N 472-ФЗ)")
    doc.add_paragraph("(часть 3.3 введена Федеральным законом от 27.12.2019 N 472-ФЗ)")
    doc.add_paragraph("Статья 53. Строительный контроль")
    doc.add_paragraph("3.3. Другая часть другой статьи.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Условие"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Высота"
    table.cell(1, 1).text = "10 м"
    doc.save(path)


def test_docx_to_structural_search_keeps_true_clause_and_edition(settings, tmp_path):
    import re
    from types import SimpleNamespace

    from qdrant_client import QdrantClient, models

    from src.common.db.qdrant_client import QdrantRepository
    from src.dvd_service.dto.fragment_search import FragmentSearchRequest
    from src.dvd_service.services.dvd_service import SearchService
    from src.dvd_service.services.fragment_search import FragmentSearchService

    path = tmp_path / "legal.docx"
    make_word_fixture(path)
    parser = DocumentParser(settings)
    raw = parser.extract_raw(str(path))
    assert any(b["text"].startswith("3.3. По решению") for b in raw)
    assert any(b["category"] == "Table" and "10 м" in b["html"] for b in raw)

    class AdverseLlm:
        def chat(self, system, user, schema):
            ids = [int(i) for i in re.findall(r"^\[(\d+)\]", user, re.M)]
            field = next(iter(schema["properties"]))
            if field == "blocks":
                return {field: [{"id": i, "boundary": "continuation"} for i in ids]}
            if field == "parts":
                return {field: [{"id": i, "merge_with_previous": True} for i in ids]}
            # Model omission/misclassification must not override explicit source anchors.
            return {
                "nodes": [
                    {
                        "id": i,
                        "type": "paragraph",
                        "numbering": "",
                        "relation": "deeper",
                        "block": "main",
                        "tags": [],
                    }
                    for i in ids
                ]
            }

    parts = parser.to_logical_parts(raw, AdverseLlm())
    StructureTagger(settings).tag(parts, AdverseLlm())
    hb = HierarchyBuilder()
    tree = hb.build(parts, StructureTagger(settings).numbering_ranks(parts))
    hb.cap_unnumbered_nesting(tree)
    hb.group_amendment(tree)
    nodes = hb.flatten(tree)
    repo = QdrantRepository.__new__(QdrantRepository)
    repo.settings = settings
    repo.collection = "regression"
    repo.client = QdrantClient(":memory:")
    repo.client.create_collection(
        repo.collection,
        vectors_config=models.VectorParams(size=2, distance=models.Distance.COSINE),
    )
    edition = "ред. от\u00a030.01.2026"
    for order, n in enumerate(nodes):
        n.update(
            order=order,
            doc_id="code",
            name="Кодекс",
            version=edition,
            versions=[edition],
        )
    repo.upsert(
        [models.PointStruct(id=n["id"], vector=[1.0, 0.0], payload=n) for n in nodes]
    )
    svc = FragmentSearchService(SearchService(repo, settings, SimpleNamespace()))
    try:
        req = FragmentSearchRequest(
            pattern="52 / 3.3", doc_id="code", version="ред. от 30.01.2026"
        )
        result = svc.search(req)
        assert result.complete and not result.ambiguous and result.match_count == 1
        assert result.hits[0].numbering == "3.3"
        assert result.hits[0].text.startswith("По решению застройщика")
        assert all(
            "Другая часть" not in h.text and "Работы выполняют" not in h.text
            for h in result.hits
        )
        assert any("часть 3.3 введена" in h.text for h in result.hits[1:])
        roots = svc.search(
            req.model_copy(update={"pattern": "3.3", "include_children": False})
        )
        assert roots.ambiguous and roots.match_count == 2
        assert all(c["excerpt"] and c["content_digest"] for c in roots.candidates)
    finally:
        repo.client.close()


def test_letter_list_items_are_siblings_under_numeric_item():
    parts = [
        {
            "id": i,
            "text": text,
            "numbering": num,
            "type": typ,
            "source_delimiter": delimiter,
        }
        for i, (text, num, typ, delimiter) in enumerate(
            [
                ("Статья 52. Строительство", "52", "article", ""),
                ("Часть", "3.3", "clause", "."),
                ("Пункт", "1", "list_item", ")"),
                ("Первое условие", "а)", "list_item", ")"),
                ("Второе условие", "б)", "list_item", ")"),
                ("Другой пункт", "2", "list_item", ")"),
            ]
        )
    ]
    hb = HierarchyBuilder()
    nodes = hb.flatten(hb.build(parts, StructureTagger(None).numbering_ranks(parts)))
    by_id = {n["id"]: n for n in nodes}
    letters = [n for n in nodes if n["numbering"] in ["а)", "б)"]]
    assert [by_id[n["parent_id"]]["numbering"] for n in letters] == ["1", "1"]
    assert by_id[nodes[-1]["parent_id"]]["numbering"] == "3.3"


@pytest.mark.parametrize(
    "text",
    [
        "Раздел 2 (Измененная редакция, Изм. № 1).",
        "Пункт 3.3 (Измененная редакция, Изм. № 1).",
        "(Измененная редакция, Изм. № 1).",
    ],
)
def test_editorial_revision_labels_are_notes_not_duplicate_headings(settings, text):
    from src.dvd_service.modules.source_structure import SourceStructure

    anchor = SourceStructure.anchor(text)
    assert anchor == {"type": "note", "numbering": "", "block": "amendment"}
    parser = DocumentParser(settings)
    assert parser._heuristic_boundary(text, "продолжение текста") == "new"
    assert parser._heuristic_boundary("предыдущий текст", text) == "new"


@pytest.mark.parametrize(
    "text", ["29  декабря  2004  года\tN 190-ФЗ", "29 декабря 2004 года N 190-ФЗ"]
)
def test_written_dates_have_no_own_number(text):
    from src.dvd_service.modules.source_structure import SourceStructure

    assert SourceStructure.anchor(text).get("numbering") == ""


def test_source_headings_ignore_inferred_sections_in_preface_and_continuations():
    from src.dvd_service.modules.source_structure import SourceStructure

    texts = [
        ("Список изменяющих документов", "section"),
        ("Глава 2. Полномочия органов государственной власти", "chapter"),
        ("ОРГАНОВ МЕСТНОГО САМОУПРАВЛЕНИЯ", "section"),
        ("Глава 6. Строительство", "chapter"),
        ("Статья 52. Осуществление строительства", "article"),
        ("3.3. По решению застройщика", "clause"),
    ]
    parts = [
        {
            "id": i,
            "text": text,
            "type": typ,
            "relation": "deeper",
            **SourceStructure.anchor(text),
        }
        for i, (text, typ) in enumerate(texts)
    ]
    hb = HierarchyBuilder()
    nodes = hb.flatten(hb.build(parts, StructureTagger(None).numbering_ranks(parts)))
    by_id = {n["id"]: n for n in nodes}
    for n in nodes:
        if n["type"] == "chapter":
            assert by_id[n["parent_id"]]["type"] == "document"
    clause = next(n for n in nodes if n["numbering"] == "3.3")
    article = by_id[clause["parent_id"]]
    assert article["numbering"] == "52"
    assert by_id[article["parent_id"]]["numbering"] == "6"


def test_source_article_keeps_chapter_after_inferred_title_resets_stack():
    from src.dvd_service.modules.source_structure import SourceStructure

    texts = [
        ("Глава 6.1. Саморегулирование", "chapter", "top"),
        ("В ОБЛАСТИ СТРОИТЕЛЬСТВА", "section", "top"),
        ("Статья 55.1. Основные цели", "article", "same"),
    ]
    parts = [
        {
            "id": i,
            "text": text,
            "type": typ,
            "relation": rel,
            **SourceStructure.anchor(text),
        }
        for i, (text, typ, rel) in enumerate(texts)
    ]
    hb = HierarchyBuilder()
    nodes = hb.flatten(hb.build(parts, StructureTagger(None).numbering_ranks(parts)))
    by_id = {n["id"]: n for n in nodes}
    article = next(n for n in nodes if n["type"] == "article")
    assert by_id[article["parent_id"]]["numbering"] == "6.1"
