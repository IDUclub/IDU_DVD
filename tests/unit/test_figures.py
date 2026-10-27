"""Figures across extraction, locked pipeline units, storage and retrieval."""

import hashlib
import io
import json
import zlib
from copy import deepcopy
from unittest.mock import Mock

import httpx
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

from src.api_clients.llm_client import OpenAICompatibleClient
from src.api_clients.vision_client import VisionClient
from src.common.auth import get_effective_user_id, require_authenticated
from src.dependencies import Dependencies
from src.dvd_service.dto import SearchRequest
from src.dvd_service.dto.fragment_search import FragmentSearchRequest
from src.dvd_service.modules.doc_parsers import DocumentParser
from src.dvd_service.modules.docx_reader import DocxReader
from src.dvd_service.modules.figures import filter_figures, image_block
from src.dvd_service.modules.pdf_reader import PdfReader, layout_blocks
from src.dvd_service.routers.library import router
from tests.unit.test_services import wired  # noqa: F401 — shared full-pipeline fixture


def png(color="blue", size=(160, 160)):
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.parametrize("before", [True, False])
@pytest.mark.parametrize(
    "caption", ["Рисунок 1 — Схема", "Рис. 2 План", "Figure 3 Plan"]
)
def test_docx_picture_and_adjacent_caption(tmp_path, before, caption):
    doc = Document()
    doc.add_paragraph("1 Область применения")
    if before:
        doc.add_paragraph(caption)
    doc.add_paragraph().add_run().add_picture(io.BytesIO(png()))
    if not before:
        doc.add_paragraph(caption)
    path = tmp_path / "figure.docx"
    doc.save(path)
    blocks = DocxReader().read(path)
    figure = next(b for b in blocks if b["category"] == "Image")
    assert figure["figure"]["caption"] == caption
    assert figure["image_bytes"].startswith(b"\x89PNG")
    assert figure["figure"]["width"] == 160


def test_docx_table_cell_and_legacy_image(tmp_path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    run = cell.paragraphs[0].add_run()
    run.add_picture(io.BytesIO(png()))
    cell.add_paragraph("Рисунок 1 — Ячейка")
    # Replace DrawingML with the legacy VML representation using the same relation.
    from docx.oxml import parse_xml

    rid = run._r.xpath(".//a:blip")[0].get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    run._r.remove(run._r.xpath(".//w:drawing")[0])
    run._r.append(
        parse_xml(
            "<w:pict "
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<v:shape><v:imagedata r:id="{rid}"/></v:shape></w:pict>'
        )
    )
    path = tmp_path / "legacy.docx"
    doc.save(path)
    blocks = DocxReader().read(path)
    assert [b["category"] for b in blocks] == ["Table", "Image"]
    assert blocks[1]["figure"]["caption"] == "Рисунок 1 — Ячейка"


def test_inline_docx_image_keeps_run_order_and_adjacent_paragraph_caption(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("До рисунка.")
    paragraph.add_run().add_picture(io.BytesIO(png()))
    paragraph.add_run("После рисунка.")
    doc.add_paragraph("Рисунок 2 — План")
    path = tmp_path / "inline.docx"
    doc.save(path)
    blocks = DocxReader().read(path)
    assert [b["category"] for b in blocks] == [
        "NarrativeText",
        "Image",
        "NarrativeText",
        "NarrativeText",
    ]
    assert blocks[0]["text"] == "До рисунка."
    assert blocks[1]["figure"]["caption"] == "Рисунок 2 — План"
    assert blocks[2]["text"] == "После рисунка."


def test_image_free_docx_hashes_and_fragments_unchanged(tmp_path, settings):
    doc = Document()
    doc.add_paragraph("Обычный текст.")
    path = tmp_path / "plain.docx"
    doc.save(path)
    expected = [
        {
            "text": "Обычный текст.",
            "category": "NarrativeText",
            "html": None,
            "page": None,
            "bbox": None,
            "source_marker": "",
            "source_paragraph": True,
        }
    ]
    parser = DocumentParser(settings)
    actual = parser.extract_raw(path)
    assert actual == expected
    assert (
        parser.content_hash(actual)
        == hashlib.sha256("Обычный текст.".encode("utf-8")).hexdigest()
    )
    assert parser.block_hashes(actual) == parser.block_hashes(expected)
    assert parser.to_logical_parts(actual, None) == parser.to_logical_parts(
        expected, None
    )


def figure_pdf(nested=False):
    """Small PDF with text, a blue raster, a vector diagram and page furniture."""

    def stream(properties, data):
        return (
            b"<< "
            + properties
            + b" /Length %d >>\nstream\n" % len(data)
            + data
            + b"\nendstream"
        )

    image = b"q 160 0 0 160 40 360 cm /Im1 Do Q"
    resource = b"/Im1 6 0 R"
    if nested:
        image = b"q 1 0 0 1 40 360 cm /Fm1 Do Q"
        resource = b"/Fm1 7 0 R"
    content = (
        b"BT /F1 12 Tf 40 560 Td "
        b"(A document with enough real text to avoid the OCR path.) Tj ET\n"
        + image
        + b"\nBT /F1 12 Tf 40 340 Td (Figure 1 Plan) Tj ET\n"
        b"0 0 0 rg 40 100 200 200 re f\n"
        b"1 1 1 RG 50 110 m 230 290 l S\n"
        b"0 0 0 RG 0 0 400 600 re S 0 40 m 400 40 l S\n"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 600] "
        b"/Resources << /Font << /F1 5 0 R >> /XObject << "
        + resource
        + b" >> >> /Contents 4 0 R >>",
        stream(b"", content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        stream(
            b"/Type /XObject /Subtype /Image /Width 4 /Height 4 "
            b"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode",
            zlib.compress(bytes([0, 0, 255]) * 16),
        ),
    ]
    if nested:
        objects.append(
            stream(
                b"/Type /XObject /Subtype /Form /BBox [0 0 160 160] "
                b"/Resources << /XObject << /Im1 6 0 R >> >>",
                b"q 160 0 0 160 0 0 cm /Im1 Do Q",
            )
        )
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


@pytest.mark.parametrize("nested", [False, True])
def test_pdf_raster_and_vector_regions(tmp_path, settings, nested):
    path = tmp_path / "figure.pdf"
    path.write_bytes(figure_pdf(nested=nested))
    blocks = DocumentParser(settings).extract_raw(path)
    images = [b for b in blocks if b["category"] == "Image"]
    assert len(images) == 2
    assert images[0]["bbox"] == [40, 80, 200, 240]
    assert images[0]["figure"]["caption"] == "Figure 1 Plan"
    assert all(
        b["page"] == 1 and b["image_bytes"].startswith(b"\x89PNG") for b in images
    )
    # The frame and thin rule are ignored; source text blocks retain their shape.
    assert images[1]["bbox"] == pytest.approx([40, 300, 240, 500])
    assert [b["category"] for b in blocks] == [
        "NarrativeText",
        "Image",
        "NarrativeText",
        "Image",
    ]


@pytest.mark.parametrize(
    "rotation,expected",
    [
        (0, [20, 50, 180, 210]),
        (90, [310, 20, 470, 180]),
        (180, [180, 310, 340, 470]),
        (270, [50, 180, 210, 340]),
    ],
)
def test_pdf_figure_bbox_matches_rotated_cropped_render(rotation, expected):
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(figure_pdf(nested=True))
    try:
        page = pdf[0]
        page.set_cropbox(20, 50, 380, 570)
        page.set_rotation(rotation)
        images = PdfReader(figure_max_pixels=4096)._page_figures(page, 1)
        raster = images[0]
        assert raster["bbox"] == expected
        with Image.open(io.BytesIO(raster["image_bytes"])) as image:
            assert image.width * image.height <= 4096
            assert image.getpixel((image.width // 2, image.height // 2)) == (0, 0, 255)
    finally:
        pdf.close()


def test_ocr_picture_crop_uses_rendered_pixels_and_pdf_points():
    blocks = layout_blocks(
        [
            {"category": "Picture", "bbox": [20, 40, 120, 140]},
            {"category": "Caption", "text": "Рис. 1 Схема"},
        ],
        2,
        png(size=(200, 200)),
        page_size=(100, 100),
    )
    image = blocks[0]
    assert image["figure"]["width"] == image["figure"]["height"] == 100
    assert image["bbox"] == [10, 20, 60, 70]
    assert image["figure"]["caption"] == "Рис. 1 Схема"
    assert (
        layout_blocks([{"category": "Picture", "bbox": [9, 9, 1, 1]}], 1, png()) == []
    )


def test_noise_filters_deduplicate_repeated_page_logos_and_cap(settings):
    settings.figure_max_per_document = 1
    unique = image_block(png("green"))
    duplicate = deepcopy(unique)
    assert filter_figures([unique, duplicate, image_block(png("red"))], settings) == [
        unique
    ]
    assert filter_figures([image_block(png(size=(4, 4)))], settings) == []
    logos = [image_block(png(), page=page) for page in range(1, 5)]
    assert filter_figures(logos, settings) == []
    assert image_block(b"invalid EMF data") is None
    large = image_block(png(size=(600, 500)), max_pixels=10_000)
    assert large["figure"]["width"] * large["figure"]["height"] <= 10_000


def test_vision_request_thinking_disabled_and_cache(tmp_path, settings):
    settings.ocr_cache_dir = str(tmp_path)
    settings.llm_think_mode = "chat_template"
    settings.llm_api_key = "fallback-key"
    settings.vision_api_key = None
    settings.vision_model = ""
    settings.vision_base_url = "http://vision.test/v1"
    calls = []

    def handler(request):
        body = json.loads(request.content)
        calls.append(body)
        assert request.headers["authorization"] == "Bearer fallback-key"
        assert body["chat_template_kwargs"] == {"enable_thinking": False}
        assert "reasoning_effort" not in body
        image = body["messages"][1]["content"][0]["image_url"]["url"]
        assert image.startswith("data:image/png;base64,")
        return httpx.Response(
            200, json={"choices": [{"message": {"content": "Схема 5 м"}}]}
        )

    client = VisionClient(settings, transport=httpx.MockTransport(handler))
    try:
        for _ in range(2):
            assert client.describe(png(), "a" * 64, "Рисунок 1") == "Схема 5 м"
    finally:
        client.close()
    assert len(calls) == 1
    assert calls[0]["model"] == settings.llm_model


@pytest.mark.parametrize(
    "mode,effort,expected",
    [
        ("reasoning_effort", "high", {"reasoning_effort": "high"}),
        ("reasoning_effort", None, {"reasoning_effort": "low"}),
        ("chat_template", "high", {"chat_template_kwargs": {"enable_thinking": True}}),
        (
            "chat_template",
            "medium",
            {"chat_template_kwargs": {"enable_thinking": True}},
        ),
        ("chat_template", "low", {"chat_template_kwargs": {"enable_thinking": False}}),
        ("chat_template", "none", {"chat_template_kwargs": {"enable_thinking": False}}),
        ("chat_template", None, {"chat_template_kwargs": {"enable_thinking": False}}),
        ("off", "high", {}),
    ],
)
def test_llm_think_mode_request_bodies(mode, effort, expected):
    seen = []

    def handler(request):
        seen.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": "{}"}}]})

    with OpenAICompatibleClient(model="gpt-oss", think_mode=mode) as client:
        client.reasoning_effort = effort
        client._client.close()
        client._client = httpx.Client(transport=httpx.MockTransport(handler))
        client.chat("system", "user", {"type": "object"})
    assert {
        k: v
        for k, v in seen[0].items()
        if k in {"reasoning_effort", "chat_template_kwargs"}
    } == expected


@pytest.mark.parametrize("mode", ["boundaries", "ranges"])
def test_pipeline_figure_payload_search_and_version_lifecycle(wired, monkeypatch, mode):
    import src.dvd_service.services.dvd_service as module

    wired.ingestion.settings.logical_partition_mode = mode
    wired.ingestion.settings.vision_enabled = True
    vision = Mock(model="vision-stub")
    vision.describe.return_value = "VISION_SENTINEL: обозначение 12 м"
    monkeypatch.setattr(module, "VisionClient", lambda _, **kwargs: vision)
    figure = image_block(png())
    figure["figure"]["caption"] = figure["text"] = "Рисунок 1 — Схема"
    raw = [{"text": "1 Область применения", "category": "NarrativeText"}, figure]
    original_hash = DocumentParser.content_hash(raw)
    result = wired.ingestion.ingest(
        "doc.docx", raw, original_hash, version_override="v1"
    )
    assert DocumentParser.content_hash(raw) == original_hash
    detail = wired.library.get_document(result["doc_id"])
    fragment = next(f for f in detail.fragments if f.kind == "figure")
    assert fragment.type == fragment.element_type == "figure"
    assert fragment.figure["description_generated"] is True
    assert fragment.figure["description_model"] == "vision-stub"
    assert fragment.image_url == f"/library/figures/{fragment.id}/image"
    assert wired.storage.download(fragment.figure["object_key"])[1] == "image/png"
    for _, prompt, schema in wired.ollama.chat_calls:
        if any(
            key in schema.get("properties", {})
            for key in ("blocks", "parts", "nodes", "fragments")
        ):
            assert "VISION_SENTINEL" not in prompt
    hits = wired.search.search(SearchRequest(query="схема", types=["figure"]), None)
    assert hits.count == 1 and hits.hits[0].image_url == fragment.image_url
    assert any(
        h.kind == "figure"
        for h in wired.search.search(SearchRequest(query="схема"), None).hits
    )
    updated = wired.ingestion.update(
        result["name"],
        "doc.docx",
        deepcopy(raw),
        original_hash + "new",
        version_override="v2",
    )
    wired.ingestion.delete_document(result["name"], "v1")
    assert any(k.startswith("figures/") for k in wired.storage.objects)
    assert updated["nodes"] > 0
    wired.ingestion.delete_document(result["name"])
    assert not wired.storage.objects


def test_disabled_or_failed_vision_keeps_captionless_figure(wired, monkeypatch):
    import src.dvd_service.services.dvd_service as module
    from src.dvd_service.modules.progress import Progress

    for enabled in (False, True):
        wired.ingestion.settings.vision_enabled = enabled
        vision = Mock()
        vision.describe.side_effect = RuntimeError("offline")
        monkeypatch.setattr(module, "VisionClient", lambda _, **kwargs: vision)
        raw = [image_block(png())]
        wired.ingestion._describe_figures(raw, Progress())
        assert raw[0]["text"] == "Рисунок без подписи"
        assert raw[0]["figure"]["description_generated"] is False


def figure_app(repo, storage, user_storage=None, user="owner"):
    app = FastAPI()
    app.include_router(router)
    for dependency, value in [
        (Dependencies.get_qdrant, repo),
        (Dependencies.get_document_storage, storage),
        (Dependencies.get_user_document_storage, user_storage or storage),
        (Dependencies.get_urban_api, Mock()),
        (get_effective_user_id, user),
        (require_authenticated, None),
    ]:

        def override(value):
            return lambda: value

        app.dependency_overrides[dependency] = override(value)
    return app


def test_figure_endpoint_png_404_and_user_scope(fake_qdrant, fake_document_storage):
    payload = {"kind": "figure", "figure": {"object_key": "figures/test.png"}}
    fake_qdrant.points["figure"] = ([1.0], payload)
    fake_qdrant.points["text"] = ([1.0], {"kind": "text"})
    fake_document_storage.upload("figures/test.png", png(), "image/png")
    with TestClient(figure_app(fake_qdrant, fake_document_storage)) as client:
        response = client.get("/library/figures/figure/image")
        assert response.status_code == 200
        assert response.headers["content-type"] == "image/png"
        assert response.content == png()
        for node_id in ("missing", "text"):
            assert client.get(f"/library/figures/{node_id}/image").status_code == 404
        payload.update(user_id="different-owner", project_id="p1")
        assert client.get("/library/figures/figure/image").status_code == 404
        payload["user_id"] = "owner"
        url = "/library/figures/figure/image"
        assert client.get(url + "?project_id=p2").status_code == 404
        assert client.get(url + "?project_id=p1").status_code == 200
        fake_document_storage.delete("figures/test.png")
        assert client.get("/library/figures/figure/image").status_code == 404


def test_figure_filtered_search_contract():
    request = FragmentSearchRequest(
        kind="figure", types=["figure"], rank_by_relevance=True, query="схема"
    )
    assert request.kind == "figure" and request.types == ["figure"]


def test_figure_endpoint_requires_authentication(fake_qdrant, fake_document_storage):
    from fastapi import HTTPException

    app = figure_app(fake_qdrant, fake_document_storage)

    def denied():
        raise HTTPException(401, "authentication required")

    app.dependency_overrides[require_authenticated] = denied
    with TestClient(app) as client:
        assert client.get("/library/figures/missing/image").status_code == 401


def test_storage_failure_cleans_up_unreferenced_objects(wired):
    wired.ingestion.settings.vision_enabled = False
    raw = [image_block(png())]
    wired.storage.fail_upload = True
    with pytest.raises(RuntimeError, match="MinIO failure"):
        wired.ingestion.ingest("doc.docx", raw, DocumentParser.content_hash(raw))
    assert not wired.storage.objects
    assert not wired.qdrant.points


def test_reload_removes_old_images_and_keeps_new_ones(wired):
    wired.ingestion.settings.vision_enabled = False
    raw = [image_block(png())]
    first = wired.ingestion.ingest("doc.docx", raw, DocumentParser.content_hash(raw))
    old_keys = set(wired.storage.objects)
    new_raw = [image_block(png("green"))]
    second = wired.ingestion.reload(
        first["name"], "doc.docx", new_raw, DocumentParser.content_hash(new_raw)
    )
    assert not old_keys.intersection(wired.storage.objects)
    assert len(wired.storage.objects) == 1
    assert second["nodes"] > 0


def test_user_pipeline_stamps_owner_and_uses_private_storage(wired):
    from src.common.db.redis_client import RedisClient
    from src.dvd_service.services.user_index_service import build_user_ingestion
    from tests.conftest import FakeDocumentStorage

    base = wired.ingestion
    base.settings.vision_enabled = False
    storage = FakeDocumentStorage()
    ingestion = build_user_ingestion(
        settings=base.settings,
        qdrant=wired.qdrant,
        redis=RedisClient(base.settings),
        storage=storage,
        jobs=wired.jobs,
        parser=base.parser,
        structure=base.structure,
        hierarchy=base.hierarchy,
        version_detector=base.version_detector,
        reference_extractor=base.reference_extractor,
        reference_resolver=base.reference_resolver,
        outbox=None,
        user_id="owner",
        project_id="p1",
    )
    raw = [image_block(png())]
    result = ingestion.ingest("user.docx", raw, DocumentParser.content_hash(raw))
    assert storage.objects and not wired.storage.objects
    point_id, payload = next(
        (i, pl) for i, (_, pl) in wired.qdrant.points.items() if pl["kind"] == "figure"
    )
    assert payload["user_id"] == "owner" and payload["project_id"] == "p1"
    app = figure_app(wired.qdrant, wired.storage, user_storage=storage)
    with TestClient(app) as client:
        assert client.get(f"/library/figures/{point_id}/image").status_code == 200
    ingestion.delete_document(result["name"])
    assert not storage.objects


def test_image_hash_changes_with_pixels_and_ignores_generated_description():
    raw = [image_block(png())]
    first = DocumentParser.content_hash(raw)
    blocks = DocumentParser.block_hashes(raw)
    raw[0]["text"] = raw[0]["figure"]["description"] = "Ответ модели"
    assert DocumentParser.content_hash(raw) == first
    assert DocumentParser.block_hashes(raw) == blocks
    assert DocumentParser.content_hash([image_block(png("green"))]) != first


def test_truncated_vision_answer_is_not_cached(tmp_path, settings):
    settings.ocr_cache_dir = str(tmp_path)
    transport = httpx.MockTransport(
        lambda _: httpx.Response(
            200,
            json={
                "choices": [
                    {"finish_reason": "length", "message": {"content": "partial"}}
                ]
            },
        )
    )
    client = VisionClient(settings, transport=transport)
    try:
        with pytest.raises(ValueError, match="truncated"):
            client.describe(png(), "b" * 64)
    finally:
        client.close()
    assert not list(tmp_path.rglob("*.json"))


def test_figure_stage_reports_progress(wired):
    wired.ingestion.settings.vision_enabled = False
    progress = Mock()
    wired.ingestion._describe_figures([image_block(png())], progress)
    progress.stage.assert_called_once_with("figures", 1)
    progress.advance.assert_called_once_with(1, 1)
    progress.complete_stage.assert_called_once()


def test_figures_keep_source_order_between_nested_clauses(settings):
    from src.dvd_service.modules.hierarchy import HierarchyBuilder

    parts = [
        {
            "id": 0,
            "text": "Section",
            "type": "section",
            "numbering": "1",
            "src_ids": [0],
        },
        {
            "id": 1,
            "text": "Clause",
            "type": "clause",
            "numbering": "1.1",
            "src_ids": [1],
        },
        {
            "id": 2,
            "text": "Рисунок без подписи",
            "type": "figure",
            "src_ids": [2],
            "figure": image_block(png())["figure"],
        },
        {
            "id": 3,
            "text": "Subclause",
            "type": "subclause",
            "numbering": "1.1.1",
            "src_ids": [3],
        },
    ]
    hierarchy = HierarchyBuilder()
    nodes = hierarchy.flatten(hierarchy.build(parts, {"1": 1, "1.1": 2, "1.1.1": 3}))
    assert [n["src_ids"] for n in nodes[1:]] == [[0], [1], [2], [3]]
    assert nodes[3]["parent_id"] == nodes[1]["id"]


def test_amendment_raw_cache_roundtrips_images(wired):
    from src.dvd_service.services.amendment_service import AmendmentService

    service = AmendmentService(wired.ingestion, Mock(), wired.ingestion.settings)
    raw = [image_block(png())]
    content_hash = DocumentParser.content_hash(raw)
    service.store_raw(content_hash, raw)
    assert service._raw("unused.docx", content_hash, "doc.docx") == raw
    assert raw[0]["image_bytes"] == png()
