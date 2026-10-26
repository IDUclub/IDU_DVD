"""Preview must match the content actually passed to Qdrant by ingestion."""

import io
import json
from unittest.mock import Mock

import pytest
from docx import Document
from PIL import Image

from scripts.preview_fragments import FragmentPreview, main
from tests.unit.test_services import wired  # noqa: F401


@pytest.mark.parametrize("mode", ["boundaries", "ranges"])
def test_preview_matches_ingestion_payload(wired, tmp_path, mode):
    service = wired.ingestion
    service.settings.logical_partition_mode = mode
    source = tmp_path / "example.docx"
    document = Document()
    document.add_heading("Статья 1. Общие положения", level=1)
    document.add_paragraph("1. Документ определяет требования к территории.")
    document.add_paragraph("2. Территория должна быть доступна для пешеходов.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    document.save(source)

    report = FragmentPreview(service.settings).build(source, wired.ollama)
    assert '"image_bytes"' not in json.dumps(report)
    assert not wired.qdrant.points
    assert not wired.ollama.embed_calls
    assert not wired.storage.objects

    raw = service.parser.extract_raw(str(source))
    service.ingest(str(source), raw, service.parser.content_hash(raw))
    payloads = sorted(
        (payload for _, payload in wired.qdrant.points.values()),
        key=lambda payload: payload["order"],
    )
    assert len(payloads) == report["fragment_count"]
    preview_ids = {node["id"]: node["order"] for node in report["fragments"]}
    stored_ids = {
        pid: payload["order"] for pid, (_, payload) in wired.qdrant.points.items()
    }
    for preview, payload in zip(report["fragments"], payloads):
        for key, value in preview.items():
            if key in {"id", "span_id", "embedding_text"}:
                continue
            if key in {"parent_id", "prev_id", "next_id"}:
                assert preview_ids.get(value) == stored_ids.get(payload[key])
            elif key in {"child_ids", "ancestor_ids"}:
                assert [preview_ids[v] for v in value] == [
                    stored_ids[v] for v in payload[key]
                ]
            else:
                assert value == payload[key], key
    embedded = [text for batch in wired.ollama.embed_calls for text in batch]
    assert embedded == [node["embedding_text"] for node in report["fragments"]]
    assert "ТАБЛИЦА (HTML)" in FragmentPreview.render(report)


@pytest.mark.parametrize("mode", ["boundaries", "ranges"])
@pytest.mark.parametrize("vision_enabled", [False, True])
def test_image_bytes_stay_out_of_preview_and_transport_payloads(
    wired, tmp_path, monkeypatch, mode, vision_enabled
):
    import src.dvd_service.services.dvd_service as module
    from src.dvd_service.dto import SearchRequest

    service = wired.ingestion
    service.settings.logical_partition_mode = mode
    service.settings.vision_enabled = vision_enabled
    vision = Mock(model="preview-vision")
    vision.describe.return_value = "Схема: высота 12 м"
    monkeypatch.setattr(module, "VisionClient", lambda _: vision)
    source = tmp_path / "figure.docx"
    png = io.BytesIO()
    Image.new("RGB", (160, 160), "blue").save(png, format="PNG")
    document = Document()
    document.add_paragraph("1 Область применения")
    document.add_paragraph().add_run().add_picture(io.BytesIO(png.getvalue()))
    document.add_paragraph("Рисунок 1 — Схема")
    document.save(source)
    report = FragmentPreview(service.settings).build(source, wired.ollama)
    preview = next(n for n in report["fragments"] if n["kind"] == "figure")
    assert not wired.storage.objects
    raw = service.parser.extract_raw(str(source))
    assert any(isinstance(b.get("image_bytes"), bytes) for b in raw)
    result = service.ingest(
        str(source),
        raw,
        service.parser.content_hash(raw),
        job_id="figure-payloads",
    )
    payloads = [payload for _, payload in wired.qdrant.points.values()]
    figure = next(p for p in payloads if p["kind"] == "figure")
    assert wired.storage.download(figure["figure"]["object_key"])[0] == png.getvalue()
    for key in (
        "text",
        "kind",
        "type",
        "element_type",
        "source_text",
        "char_start",
        "char_end",
    ):
        assert preview[key] == figure[key], key
    assert preview["figure"] == {
        k: v for k, v in figure["figure"].items() if k != "object_key"
    }
    search = wired.search.search(SearchRequest(query="схема"), None).model_dump()
    library = wired.library.get_document(result["doc_id"]).model_dump()
    job = wired.jobs.get("figure-payloads")
    event = wired.outbox.peek()
    assert job["status"] == "done" and event is not None
    for artifact in (report, payloads, search, library, job, event):
        encoded = json.dumps(artifact)
        assert '"image_bytes"' not in encoded
        assert '"image_base64"' not in encoded


def test_invalid_source_fails_before_loading_llm(tmp_path):
    with pytest.raises(SystemExit) as exc:
        main([str(tmp_path / "missing.docx")])
    assert exc.value.code == 2


def test_cli_writes_reports_and_closes_client(wired, tmp_path, monkeypatch):
    import src.api_clients

    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("1. Требования к доступности территории.")
    document.save(source)
    output = tmp_path / "report"
    monkeypatch.setattr(src.api_clients, "create_llm", lambda: wired.ollama)
    assert main([str(source), "--output-dir", str(output)]) == 0
    report = json.loads((output / "fragments.json").read_text(encoding="utf-8"))
    assert report["fragment_count"] > 0
    assert "ТЕКСТ ФРАГМЕНТА" in (output / "fragments.txt").read_text(encoding="utf-8")
    assert wired.ollama.closed
    assert not wired.ollama.embed_calls
    assert not wired.qdrant.points
