"""Word counters are read from OOXML, including overrides and inherited styles."""

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from src.dvd_service.modules.docx_reader import DocxReader


def test_override_restart_style_inheritance_and_per_document_state(tmp_path):
    doc = Document()
    definitions = doc.part.numbering_part.element
    definitions.append(
        parse_xml(f"""<w:abstractNum {nsdecls('w')} w:abstractNumId="900">
      <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
      <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
    </w:abstractNum>""")
    )
    definitions.append(
        parse_xml(f"""<w:num {nsdecls('w')} w:numId="900"><w:abstractNumId w:val="900"/>
      <w:lvlOverride w:ilvl="0"><w:startOverride w:val="3"/></w:lvlOverride></w:num>""")
    )
    style = doc.styles.add_style("LegalCounter", WD_STYLE_TYPE.PARAGRAPH)
    num = style.element.get_or_add_pPr().get_or_add_numPr()
    num.get_or_add_numId().val = 900
    num.get_or_add_ilvl().val = 0
    derived = doc.styles.add_style("DerivedCounter", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = style
    for text, level in [
        ("Parent", 0),
        ("Child", 1),
        ("Second child", 1),
        ("Next parent", 0),
        ("Restarted child", 1),
    ]:
        p = doc.add_paragraph(text, style=derived)
        if level:
            p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_ilvl().val = level
    path = tmp_path / "counters.docx"
    doc.save(path)
    reader = DocxReader()
    expected = [
        "3. Parent",
        "3.1. Child",
        "3.2. Second child",
        "4. Next parent",
        "4.1. Restarted child",
    ]
    assert [b["text"] for b in reader.read(path)] == expected
    assert [b["text"] for b in reader.read(path)] == expected


def test_tables_preserve_merged_cells_and_escape_html(tmp_path):
    doc = Document()
    doc.add_paragraph("Before")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "A < B"
    table.cell(0, 1).text = "First"
    table.cell(1, 1).text = "Second"
    doc.add_paragraph("After")
    path = tmp_path / "table.docx"
    doc.save(path)
    raw = DocxReader().read(path)
    assert [b["category"] for b in raw] == ["NarrativeText", "Table", "NarrativeText"]
    assert raw[1]["text"].count("A < B") == 1
    assert 'rowspan="2"' in raw[1]["html"] and "A &lt; B" in raw[1]["html"]
