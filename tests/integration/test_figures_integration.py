"""DOCX figure through the real pipeline, Qdrant/Redis/MinIO and HTTP image proxy."""

import io
import uuid

import pytest
import urllib3
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from minio import Minio
from PIL import Image

from src.common.auth import get_effective_user_id, require_authenticated
from src.common.db.minio_client import DocumentStorage
from src.common.db.qdrant_client import QdrantRepository
from src.common.db.redis_client import DocumentRegistry, JobStore, UserIndexRegistry
from src.dependencies import Dependencies
from src.dvd_service.dto import SearchRequest
from src.dvd_service.modules.doc_parsers import DocumentParser
from src.dvd_service.modules.hierarchy import HierarchyBuilder
from src.dvd_service.modules.references import ReferenceExtractor, ReferenceResolver
from src.dvd_service.modules.structure import StructureTagger
from src.dvd_service.modules.tagging import VersionDetector
from src.dvd_service.routers.library import router
from src.dvd_service.services import dvd_service as module

pytestmark = pytest.mark.integration


@pytest.fixture
def figure_storage(live_settings):
    client = Minio(
        live_settings.minio_endpoint,
        access_key=live_settings.minio_access_key,
        secret_key=live_settings.minio_secret_key,
        secure=live_settings.minio_secure,
        http_client=urllib3.PoolManager(
            timeout=urllib3.Timeout(connect=2, read=3), retries=False
        ),
    )
    storage = DocumentStorage(client, "itest-figures-" + uuid.uuid4().hex[:12])
    try:
        storage.ensure_bucket()
    except Exception as exc:  # noqa: BLE001
        pytest.skip(f"MinIO unavailable on the local stack: {exc}")
    yield storage
    for obj in client.list_objects(storage.bucket, recursive=True):
        client.remove_object(storage.bucket, obj.object_name)
    client.remove_bucket(storage.bucket)


def test_docx_figure_ingestion_search_image_and_delete(
    temp_collection,
    require_redis,
    require_ollama,
    require_embedder,
    figure_storage,
    tmp_path,
    monkeypatch,
):
    settings = temp_collection.model_copy(
        update={
            "vision_enabled": True,
            "enable_reference_linking": False,
            "enable_relations": False,
        }
    )
    settings.vector_size = len(require_embedder.embed_query("figure dimension probe"))
    monkeypatch.setattr(module, "create_llm", lambda: require_ollama)

    class VisionStub:
        model = "integration-vision-stub"

        def __init__(self, _):
            pass

        def describe(self, data, sha256, context):
            assert data.startswith(b"\x89PNG") and "Рисунок 1" in context
            return "Схема водоснабжения, труба диаметром 50 мм"

        def close(self):
            pass

    monkeypatch.setattr(module, "VisionClient", VisionStub)
    buffer = io.BytesIO()
    Image.new("RGB", (200, 160), "blue").save(buffer, format="PNG")
    doc = Document()
    doc.add_paragraph("СП 99.99999.2099 Тестовый документ")
    doc.add_paragraph("1 Область применения")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(buffer.getvalue()))
    doc.add_paragraph("Рисунок 1 — Схема водоснабжения")
    path = tmp_path / "figure.docx"
    doc.save(path)
    parser = DocumentParser(settings)
    qdrant = QdrantRepository(settings)
    qdrant.ensure_collection()
    registry = DocumentRegistry(require_redis, prefix=settings.registry_prefix)
    service = module.IngestionService(
        parser,
        StructureTagger(settings),
        HierarchyBuilder(),
        VersionDetector(),
        ReferenceExtractor(settings),
        ReferenceResolver(qdrant, registry, settings),
        qdrant,
        registry,
        figure_storage,
        JobStore(require_redis),
        settings,
    )
    raw = parser.extract_raw(path)
    result = service.ingest(
        str(path),
        raw,
        parser.content_hash(raw),
        name_override="ICII-205 integration",
        version_override="v1",
    )
    try:
        search = module.SearchService(
            qdrant,
            settings,
            UserIndexRegistry(require_redis, prefix=settings.registry_prefix),
        )
        hits = search.search(
            SearchRequest(query="водоснабжение 50 мм", types=["figure"]), None
        )
        assert hits.count == 1
        figure = hits.hits[0]
        assert figure.kind == figure.element_type == "figure"
        assert figure.figure["description_model"] == "integration-vision-stub"
        assert not figure.image_url.startswith("http")
        detail = module.LibraryService(qdrant, registry).get_document(result["doc_id"])
        assert any(f.kind == "figure" for f in detail.fragments)
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[require_authenticated] = lambda: None
        app.dependency_overrides[get_effective_user_id] = lambda: "integration-owner"
        app.dependency_overrides[Dependencies.get_qdrant] = lambda: qdrant
        for getter in (
            Dependencies.get_document_storage,
            Dependencies.get_user_document_storage,
        ):
            app.dependency_overrides[getter] = lambda: figure_storage
        app.dependency_overrides[Dependencies.get_urban_api] = lambda: None
        with TestClient(app) as client:
            response = client.get(figure.image_url)
            assert response.status_code == 200
            assert response.headers["content-type"] == "image/png"
            assert response.content.startswith(b"\x89PNG")
    finally:
        service.delete_document(result["name"])
    assert not list(
        figure_storage.client.list_objects(figure_storage.bucket, recursive=True)
    )
